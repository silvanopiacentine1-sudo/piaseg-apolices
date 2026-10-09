"""Painel de apólices canceladas (/canceladas/...). Mesmo fluxo do painel de não renovadas, com o mês
dado pela data do cancelamento. Login, vínculos e gestores são os mesmos."""

import io
from collections import defaultdict
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import extract
from sqlalchemy.orm import Session

import models
from acesso import eh_admin, exigir_admin, franqueados_do_usuario
from auth import UsuarioAtual
from database import get_db
from importer_canceladas import ler_relatorio

router = APIRouter(prefix="/canceladas")

STATUS_OPCOES = [
    "Cliente refez o seguro com a Piaseg (nova apólice/seguradora)",
    "Em negociação para recuperar o cliente",
    "Cliente vendeu o bem / encerrou a atividade",
    "Falta de pagamento (inadimplência)",
    "Cliente achou caro / reduziu custos",
    "Migrou para outra corretora",
    "Insatisfação com a seguradora ou atendimento",
    "Cancelado pela seguradora",
    "Sem retorno do cliente",
    "Outro motivo",
]

Apolice = models.CancApolice


def escopo(db: Session, user: dict, franqueado: Optional[str]):
    q = db.query(Apolice)
    if eh_admin(db, user["usuario"]):
        if franqueado:
            q = q.filter(Apolice.franqueado == franqueado)
        return q
    meus = franqueados_do_usuario(db, user["usuario"])
    if franqueado:
        if franqueado not in meus:
            raise HTTPException(403, "Você não tem acesso a este franqueado.")
        meus = [franqueado]
    return q.filter(Apolice.franqueado.in_(meus or ["__nenhum__"]))


def apolice_json(a) -> dict:
    return {
        "id": a.id,
        "franqueado": a.franqueado,
        "unidade": a.unidade,
        "cliente": a.cliente,
        "seguradora": a.seguradora,
        "cancelamento": a.cancelamento.isoformat(),
        "vigencia_inicio": a.vigencia_inicio.isoformat() if a.vigencia_inicio else None,
        "vigencia_fim": a.vigencia_fim.isoformat() if a.vigencia_fim else None,
        "numero": a.numero,
        "produto": a.produto,
        "premio": a.premio,
        "status": a.status,
        "respondido_em": a.respondido_em.isoformat() + "Z" if a.respondido_em else None,
        "no_ultimo_relatorio": a.no_ultimo_relatorio,
    }


def filtrar(q, ano: Optional[int], mes: Optional[int], situacao: Optional[str], busca: Optional[str]):
    if ano:
        q = q.filter(extract("year", Apolice.cancelamento) == ano)
    if mes:
        q = q.filter(extract("month", Apolice.cancelamento) == mes)
    if situacao == "pendente":
        q = q.filter(Apolice.status.is_(None))
    elif situacao == "respondida":
        q = q.filter(Apolice.status.isnot(None))
    if busca:
        termo = f"%{busca.strip()}%"
        q = q.filter(
            Apolice.cliente.ilike(termo)
            | Apolice.numero.ilike(termo)
            | Apolice.produto.ilike(termo)
            | Apolice.seguradora.ilike(termo)
        )
    return q


@router.get("/status-opcoes")
def status_opcoes():
    return STATUS_OPCOES


@router.get("/anos")
def anos(user: dict = UsuarioAtual, db: Session = Depends(get_db)):
    q = escopo(db, user, None).with_entities(extract("year", Apolice.cancelamento)).distinct()
    lista = sorted({int(a[0]) for a in q}, reverse=True)
    return lista or [date.today().year]


@router.get("/resumo")
def resumo(ano: int, franqueado: Optional[str] = None, user: dict = UsuarioAtual, db: Session = Depends(get_db)):
    q = filtrar(escopo(db, user, franqueado), ano, None, None, None)
    meses = {m: {"mes": m, "qtd": 0, "premio": 0.0, "respondidas": 0} for m in range(1, 13)}
    por_franqueado = defaultdict(lambda: {"qtd": 0, "premio": 0.0, "respondidas": 0, "meses": defaultdict(int)})
    por_status = defaultdict(int)
    for a in q:
        m = meses[a.cancelamento.month]
        m["qtd"] += 1
        m["premio"] += a.premio or 0
        f = por_franqueado[a.franqueado]
        f["qtd"] += 1
        f["premio"] += a.premio or 0
        f["meses"][a.cancelamento.month] += 1
        if a.status:
            m["respondidas"] += 1
            f["respondidas"] += 1
            por_status[a.status] += 1
    total = sum(m["qtd"] for m in meses.values())
    respondidas = sum(m["respondidas"] for m in meses.values())
    return {
        "total": total,
        "premio": round(sum(m["premio"] for m in meses.values()), 2),
        "respondidas": respondidas,
        "pendentes": total - respondidas,
        "meses": [dict(m, premio=round(m["premio"], 2)) for m in meses.values()],
        "por_status": [{"status": s, "qtd": por_status.get(s, 0)} for s in STATUS_OPCOES if por_status.get(s)],
        "franqueados": sorted(
            (
                {
                    "franqueado": nome,
                    "qtd": f["qtd"],
                    "premio": round(f["premio"], 2),
                    "respondidas": f["respondidas"],
                    "meses": {str(k): v for k, v in f["meses"].items()},
                }
                for nome, f in por_franqueado.items()
            ),
            key=lambda x: -x["qtd"],
        ),
    }


@router.get("/apolices")
def listar(
    ano: Optional[int] = None,
    mes: Optional[int] = None,
    franqueado: Optional[str] = None,
    situacao: Optional[str] = None,
    busca: Optional[str] = None,
    user: dict = UsuarioAtual,
    db: Session = Depends(get_db),
):
    q = filtrar(escopo(db, user, franqueado), ano, mes, situacao, busca)
    return [apolice_json(a) for a in q.order_by(Apolice.cancelamento, Apolice.cliente)]


def _apolice_permitida(db: Session, user: dict, apolice_id: int):
    a = escopo(db, user, None).filter(Apolice.id == apolice_id).first()
    if not a:
        raise HTTPException(404, "Apólice não encontrada.")
    return a


@router.get("/apolices/{apolice_id}/respostas")
def respostas(apolice_id: int, user: dict = UsuarioAtual, db: Session = Depends(get_db)):
    _apolice_permitida(db, user, apolice_id)
    rs = db.query(models.CancResposta).filter(models.CancResposta.apolice_id == apolice_id).order_by(models.CancResposta.criado_em.desc())
    return [
        {"id": r.id, "status": r.status, "comentario": r.comentario, "autor": r.autor_nome or r.autor_usuario,
         "criado_em": r.criado_em.isoformat() + "Z"}
        for r in rs
    ]


class RespostaIn(BaseModel):
    status: str
    comentario: str = ""


@router.post("/apolices/{apolice_id}/respostas")
def responder(apolice_id: int, dados: RespostaIn, user: dict = UsuarioAtual, db: Session = Depends(get_db)):
    if dados.status not in STATUS_OPCOES:
        raise HTTPException(400, "Escolha uma das opções de motivo.")
    if dados.status == "Outro motivo" and not dados.comentario.strip():
        raise HTTPException(400, "Descreva o motivo no comentário.")
    a = _apolice_permitida(db, user, apolice_id)
    agora = datetime.utcnow()
    db.add(models.CancResposta(apolice_id=a.id, status=dados.status, comentario=dados.comentario.strip(),
                               autor_usuario=user["usuario"], autor_nome=user["nome"], criado_em=agora))
    a.status = dados.status
    a.respondido_em = agora
    db.commit()
    return apolice_json(a)


@router.get("/exportar")
def exportar(
    ano: Optional[int] = None,
    mes: Optional[int] = None,
    franqueado: Optional[str] = None,
    situacao: Optional[str] = None,
    busca: Optional[str] = None,
    user: dict = UsuarioAtual,
    db: Session = Depends(get_db),
):
    q = filtrar(escopo(db, user, franqueado), ano, mes, situacao, busca)
    ultimas = {}
    for r in db.query(models.CancResposta).order_by(models.CancResposta.criado_em):
        ultimas[r.apolice_id] = r
    wb = Workbook()
    ws = wb.active
    ws.title = "Canceladas"
    ws.append(["Franqueado", "Cliente", "Seguradora", "Ramo", "Apólice", "Cancelamento", "Início vigência", "Fim vigência",
               "Prêmio líquido", "Motivo", "Comentário", "Respondido por", "Respondido em"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="072A3C")
    for a in q.order_by(Apolice.franqueado, Apolice.cancelamento):
        r = ultimas.get(a.id)
        ws.append([a.franqueado, a.cliente, a.seguradora, a.produto, a.numero, a.cancelamento, a.vigencia_inicio,
                   a.vigencia_fim, a.premio, a.status or "Pendente", r.comentario if r else "",
                   (r.autor_nome or r.autor_usuario) if r else "", r.criado_em if r else None])
    for col, larg in zip("ABCDEFGHIJKLM", [28, 40, 20, 18, 18, 13, 13, 13, 14, 44, 50, 24, 18]):
        ws.column_dimensions[col].width = larg
    for linha in ws.iter_rows(min_row=2):
        for i in (5, 6, 7):
            linha[i].number_format = "DD/MM/YYYY"
        linha[8].number_format = '"R$" #,##0.00'
        linha[12].number_format = "DD/MM/YYYY HH:MM"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="apolices-canceladas.xlsx"'},
    )


@router.post("/admin/importar")
async def importar(arquivo: UploadFile = File(...), user: dict = Depends(exigir_admin), db: Session = Depends(get_db)):
    conteudo = await arquivo.read()
    try:
        registros, ignoradas = ler_relatorio(conteudo)
    except ValueError as e:
        raise HTTPException(400, str(e))
    except Exception:
        raise HTTPException(400, "Não consegui ler o arquivo. Envie o relatório de cancelamentos do Quiver em .XLS ou .XLSX.")
    if not registros:
        raise HTTPException(400, "Nenhuma apólice encontrada na planilha.")

    existentes = {a.chave: a for a in db.query(Apolice)}
    novas = atualizadas = 0
    agora = datetime.utcnow()
    chaves = {r["chave"] for r in registros}
    for reg in registros:
        a = existentes.get(reg["chave"])
        if a is None:
            db.add(Apolice(**reg, importado_em=agora, no_ultimo_relatorio=True))
            novas += 1
        else:
            for campo, valor in reg.items():
                setattr(a, campo, valor)
            a.no_ultimo_relatorio = True
            atualizadas += 1
    # Apólices no mesmo período que não vieram no relatório (ex.: cancelamento estornado).
    # Nunca apagamos (as respostas dos franqueados ficam preservadas), só marcamos.
    inicio = min(r["cancelamento"] for r in registros)
    fim = max(r["cancelamento"] for r in registros)
    for a in existentes.values():
        if a.chave not in chaves and inicio <= a.cancelamento <= fim:
            a.no_ultimo_relatorio = False
    db.add(models.CancImportacao(arquivo=arquivo.filename or "", autor=user["nome"], linhas=len(registros),
                                 novas=novas, atualizadas=atualizadas, ignoradas=ignoradas, criado_em=agora))
    db.commit()
    return {"linhas": len(registros), "novas": novas, "atualizadas": atualizadas, "ignoradas": ignoradas,
            "periodo": [inicio.isoformat(), fim.isoformat()]}


@router.get("/admin/importacoes")
def importacoes(user: dict = Depends(exigir_admin), db: Session = Depends(get_db)):
    return [
        {"id": i.id, "arquivo": i.arquivo, "autor": i.autor, "linhas": i.linhas, "novas": i.novas,
         "atualizadas": i.atualizadas, "ignoradas": i.ignoradas, "criado_em": i.criado_em.isoformat() + "Z"}
        for i in db.query(models.CancImportacao).order_by(models.CancImportacao.criado_em.desc()).limit(30)
    ]
